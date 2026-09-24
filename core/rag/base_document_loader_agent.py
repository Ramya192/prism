# core/rag/base_document_loader_agent.py
# Shared chunk/embed/store mechanics for every document-shaped domain's
# loader agent. Extracted from domains/banking/documents/agents/
# document_loader_agent.py and domains/payroll_hr/payroll/agents/document_loader_agent.py
# once a second, near-identical copy made the duplication undeniable —
# payroll's own file docstring already said "same chunk/embed/store
# mechanics (that part of the pipeline has nothing banking-specific in it)."
# A third RAG domain (payroll_hr's planned handbook RAG) would have
# been a third copy; refactored to a shared base before that happened.
#
# Domain-specific behavior is a single override point (_extract_metadata),
# not a config flag — banking documents' bank-name detection becomes a
# one-method override in its subclass instead of an if/else branch here.
#
# Any-format ingestion (docs/UNIFIED_INGESTION_VISION.md): load() accepts
# PDF, DOCX, or CSV, decided by file extension, never by which capability
# the caller wants -- format only changes HOW text is extracted, never
# WHETHER RAG or fraud-scoring runs. CSV is fundamentally different from
# the other two: there's no continuous prose to chunk, so each row
# becomes its own retrievable unit (a rendered "col: val, ..." line) AND
# is returned as a raw dict in `records` -- the one piece of information
# genuinely new here, since a CSV row is also a fraud-detector-ready
# record, unlike a PDF/DOCX chunk of prose. PDF/DOCX return `records=None`;
# mapping their extracted text into a fraud-detector record happens one
# level up (UnifiedDomainPipeline), via the LLM extraction step, not here.
#
# IMPORTANT: banking documents is live in production and already RAGAS-
# evaluated. The PDF path here changes zero logic and zero prompt/
# behavior for it — every PDF-related line is the exact same code that
# lived in its document_loader_agent.py, just parameterized by a
# `settings` argument instead of importing one fixed Settings class
# directly.

from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_openai import OpenAIEmbeddings
from pathlib import Path
import hashlib
import pdfplumber
import docx
import pandas as pd

# CSV rows above this count still get every row scored by the fraud
# detector (that's the whole point), but only this many are embedded as
# chattable chunks -- embedding tens of thousands of rows for RAG chat
# has real API cost and no realistic user asks "tell me about row 40,000"
# instead of a targeted question retrieval already answers well.
MAX_CSV_CHUNK_ROWS = 500


class DuplicateDocumentError(ValueError):
    """Raised when an upload's exact bytes were already ingested under
    some document name -- see BaseDocumentLoaderAgent.find_duplicate()."""

    def __init__(self, filename: str, existing: str):
        self.filename, self.existing = filename, existing
        super().__init__(
            f"'{filename}' was already ingested (identical content stored as '{display_name(existing)}') "
            "-- not ingesting it again."
        )


class EmptyDocumentError(ValueError):
    """Raised when a file yields no extractable text -- typically a scanned/
    image-only PDF (there is no OCR step) or an empty DOCX/CSV. Nothing can
    be chunked, embedded, retrieved or scored from it, and Chroma rejects an
    empty add outright, so this is caught up front with a message the UI can
    show instead of a raw stack trace."""

    def __init__(self, filename: str):
        self.filename = filename
        super().__init__(
            f"'{filename}' has no extractable text -- it may be a scanned/image-only "
            "document (OCR isn't supported) or an empty file."
        )


# A session-owned document is stored as "<filename>@<owner>" so two visitors
# uploading a same-named file never overwrite each other's chunks (the
# public deployment shares one vector store across everyone), and
# display_name() strips the suffix again for anything user-facing.
OWNER_SEP = "@"


def stored_name(filename: str, owner: str | None) -> str:
    return f"{filename}{OWNER_SEP}{owner}" if owner else filename


def display_name(stored: str) -> str:
    """Inverse of stored_name(), for a name that may or may not carry an
    owner suffix (reference-corpus and test documents never do)."""
    head, sep, tail = stored.rpartition(OWNER_SEP)
    # An owner id is alphanumeric; a filename's tail ("b.pdf" in "a@b.pdf")
    # has a dot, so an '@' inside an unowned filename isn't mistaken for one.
    return head if sep and head and tail.isalnum() else stored


class BaseDocumentLoaderAgent:
    def __init__(self, settings):
        self.settings = settings
        self.splitter = RecursiveCharacterTextSplitter(
            chunk_size=settings.CHUNK_SIZE,
            chunk_overlap=settings.CHUNK_OVERLAP,
        )
        self.embedder = OpenAIEmbeddings(model=settings.EMBEDDING_MODEL)
        self._init_vector_db()

    def _extract_metadata(self, full_text: str) -> dict:
        """Override for domain-specific per-document metadata (e.g.
        banking documents' bank_name detection). Empty by default -- most
        document-shaped domains (payroll included) don't need this."""
        return {}

    @staticmethod
    def _extract_pdf_text(file_path: str) -> str:
        full_text = ""
        with pdfplumber.open(file_path) as pdf:
            for page in pdf.pages:
                full_text += page.extract_text() or ""
        return full_text

    @staticmethod
    def _extract_docx_text(file_path: str) -> str:
        d = docx.Document(file_path)
        parts = [p.text for p in d.paragraphs if p.text.strip()]
        for table in d.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(parts)

    @staticmethod
    def _extract_csv(file_path: str) -> tuple[list[str], list[dict]]:
        """Returns (chunks, records). Each row becomes one chattable
        chunk (a rendered "col: val, ..." line, not the raw CSV text --
        raw CSV text chunked by character count would slice rows apart
        arbitrarily) and one raw dict for fraud scoring. Every row is
        included in `records`; only the first MAX_CSV_CHUNK_ROWS are
        embedded as chunks (see the module-level comment for why)."""
        try:
            df = pd.read_csv(file_path)
        except pd.errors.EmptyDataError:
            return [], []   # load() turns "no chunks" into EmptyDocumentError
        records = df.to_dict(orient="records")
        chunk_rows = records[:MAX_CSV_CHUNK_ROWS]
        chunks = [
            ", ".join(f"{k}: {v}" for k, v in row.items())
            for row in chunk_rows
        ]
        return chunks, records

    def _init_vector_db(self):
        from chromadb import PersistentClient

        self.chroma_client = PersistentClient(path=self.settings.CHROMA_PATH)
        self.collection = self.chroma_client.get_or_create_collection(name=self.settings.COLLECTION_NAME)

    @staticmethod
    def _content_hash(file_path: str) -> str:
        return hashlib.sha256(Path(file_path).read_bytes()).hexdigest()

    def find_duplicate(self, file_path: str, owner: str | None = None) -> str | None:
        """Name of an already-ingested document whose file bytes are
        identical to `file_path`'s, else None. Keyed on a content hash
        stored in every chunk's metadata (not the filename -- uploads
        arrive as temp files, and a renamed copy is still a duplicate),
        so it holds across batches and for as long as the vector store
        does. Documents ingested before this check existed carry no hash
        and are never reported as duplicates.

        `owner` (a browser-session id) scopes the check to that owner's
        own documents: on the shared public store, another visitor having
        uploaded the same file must neither block this one nor leak the
        other's filename. None checks the whole store (CLI/tests)."""
        content_hash = self._content_hash(file_path)
        where = {"content_hash": content_hash}
        if owner:
            where = {"$and": [where, {"owner": owner}]}
        found = self.collection.get(where=where, limit=1)
        return found["metadatas"][0]["source"] if found["ids"] else None

    def load(self, file_path: str, filename: str = None, owner: str | None = None) -> dict:
        """Extension decides HOW text is extracted; it never decides
        whether RAG or fraud-scoring runs (both always do, one level up
        in UnifiedDomainPipeline) -- see the module docstring. `owner`
        namespaces the stored document (see stored_name()); None keeps
        the plain filename, exactly as before."""
        filename = filename or Path(file_path).stem
        ext = Path(file_path).suffix.lower()

        records = None
        if ext == ".csv":
            chunks, records = self._extract_csv(file_path)
            full_text = "\n".join(chunks)
        elif ext == ".docx":
            full_text = self._extract_docx_text(file_path)
            chunks = self.splitter.split_text(full_text)
        elif ext in (".txt", ".md"):
            # Plain text -- used for curated static reference-corpus
            # documents (see core/rag/reference_corpus.py), not typically
            # a user's own upload.
            full_text = Path(file_path).read_text(encoding="utf-8")
            chunks = self.splitter.split_text(full_text)
        else:
            full_text = self._extract_pdf_text(file_path)
            chunks = self.splitter.split_text(full_text)

        if not chunks:
            raise EmptyDocumentError(filename)

        filename = stored_name(filename, owner)
        extra_metadata = {**self._extract_metadata(full_text), "content_hash": self._content_hash(file_path)}
        if owner:
            extra_metadata["owner"] = owner
        vectors = self.embedder.embed_documents(chunks)

        self._store_chroma(filename, chunks, vectors, extra_metadata)
        return {
            "status": "success",
            "document": filename,
            "chunks_stored": len(chunks),
            "file_type": ext.lstrip("."),
            "records": records,   # list[dict] for CSV, None for PDF/DOCX
        }

    def _store_chroma(self, filename: str, chunks: list, vectors: list, extra_metadata: dict):
        # Remove duplicates
        existing = self.collection.get(where={"source": filename})
        if existing["ids"]:
            self.collection.delete(ids=existing["ids"])

        self.collection.add(
            documents=chunks,
            embeddings=vectors,
            ids=[f"{filename}_chunk_{i}" for i in range(len(chunks))],
            metadatas=[{"source": filename, "chunk_index": i, **extra_metadata} for i in range(len(chunks))],
        )
